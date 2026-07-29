# -*- coding: utf-8 -*-
"""Build Word docx for 3-ch LVDS image compare tech disclosure with embedded figures."""
import os
import re
from docx import Document
from docx.shared import Pt, Cm, Inches
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING

BASE = os.path.dirname(os.path.abspath(__file__))
MD = os.path.join(BASE, "技术交底书-3通道LVDS图像输入比较.md")
OUT = os.path.join(BASE, "技术交底书-3通道LVDS图像输入比较.docx")
FIG = os.path.join(BASE, "figures")

FIG_MAP = {
    "figures/fig1_overall.png": "fig1_overall.png",
    "figures/fig2_boards.png": "fig2_boards.png",
    "figures/fig3_fpga.png": "fig3_fpga.png",
    "figures/fig4_lvds_ch.png": "fig4_lvds_ch.png",
    "figures/fig5_modes.png": "fig5_modes.png",
    "figures/fig6_datapath.png": "fig6_datapath.png",
}


def set_run_font(run, size=12, bold=False, east="宋体"):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = east
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), east)


def add_text_para(doc, text, size=12, bold=False, center=False, first_indent=False,
                  space_before=0, space_after=6, east="宋体"):
    p = doc.add_paragraph()
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    if first_indent:
        pf.first_line_indent = Cm(0.74)
    run = p.add_run(text)
    set_run_font(run, size, bold, east)
    return p


def add_image(doc, filename, width_inches=6.2):
    path = os.path.join(FIG, filename)
    if not os.path.exists(path):
        add_text_para(doc, "[缺图: %s]" % filename, size=10, east="楷体")
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run()
    run.add_picture(path, width=Inches(width_inches))


def is_section(line):
    return bool(re.match(r"^[一二三四五六七八九十]+、", line))


def is_subsection(line):
    return line.startswith("（") and ("）" in line[:8])


SUBTITLES = {
    "基于 PC 工控机 + 采集卡的离线比对方案",
    "基于纯软件帧缓存比对的嵌入式方案",
    "基于 FPGA 的单通道或时分复用比对方案",
    "现有技术不足：",
    "以太网接口",
    "以太网协议栈与命令服务",
    "Zynq UltraScale+ MPSoC（FPGA+ARM）",
    "DDR 存储器",
    "配置与启动",
    "时钟",
    "解串芯片",
    "远端串行（系统配套，可选）",
    "LVDS/OLDI 接口",
    "I2C 与 GPIO",
    "电源模块",
    "机械结构",
    "参考帧捕获与帧差使能",
    "三路并行在线比对",
    "ROI 目标色统计与点采样",
    "解串前端配置与复位",
    "视频监视与推流",
    "错误定位与回归",
}

SHORTCOMING_PREFIXES = (
    "主机依赖重",
    "多通道并行能力不足",
    "与 GMSL/OLDI 链路耦合弱",
    "比较策略封闭",
    "部署与接口不统一",
    "扩展与维护成本高",
    "主要器件选型",
)

CAPTION_KEYS = (
    "图像输入比较",
    "载板",
    "Capture",
    "SoC",
    "FPGA",
    "LVDS",
    "OLDI",
    "工作方式",
    "数据路径",
    "接收与比较",
)


def main():
    with open(MD, "r", encoding="utf-8") as f:
        lines = f.read().splitlines()

    doc = Document()
    for sec in doc.sections:
        sec.top_margin = Cm(2.54)
        sec.bottom_margin = Cm(2.54)
        sec.left_margin = Cm(2.8)
        sec.right_margin = Cm(2.8)

    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            i += 1
            continue

        m = re.match(r"!\[.*?\]\((.*?)\)", line)
        if m:
            rel = m.group(1).replace("\\", "/")
            name = FIG_MAP.get(rel, os.path.basename(rel))
            add_image(doc, name)
            i += 1
            continue

        if line.startswith("（以下为内部"):
            add_text_para(doc, line, size=9, space_before=14, east="楷体")
            i += 1
            while i < len(lines):
                if lines[i].strip():
                    add_text_para(doc, lines[i].strip(), size=9, east="楷体")
                i += 1
            break

        if line == "产品技术交底书":
            add_text_para(doc, line, size=22, bold=True, center=True, space_after=18, east="黑体")
        elif is_section(line) or line.startswith("说明："):
            add_text_para(doc, line, size=14, bold=True, space_before=14, space_after=8, east="黑体")
        elif is_subsection(line):
            add_text_para(doc, line, size=12, bold=True, space_before=10, space_after=4, east="黑体")
        elif line in SUBTITLES or line.startswith(SHORTCOMING_PREFIXES):
            add_text_para(doc, line, size=12, bold=True, space_before=8, space_after=2, east="黑体")
        elif line.startswith("图") and any(k in line for k in CAPTION_KEYS):
            add_text_para(doc, line, size=10.5, center=True, space_after=10, east="楷体")
        elif re.match(r"^\d+\.\s", line):
            add_text_para(doc, line, size=12, space_before=6, space_after=2)
        else:
            add_text_para(doc, line, size=12, first_indent=True, space_after=4)

        i += 1

    doc.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    main()
